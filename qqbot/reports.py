"""Application runtime exporter. The ledger is authoritative; XLSX is a snapshot."""
import os
import json
import math
import threading
import uuid
from datetime import date, datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .submissions import FIELD_NAMES, REQUIRED, friendly_reason


def local_datetime(value):
    if value is None:
        return None
    return datetime.fromisoformat(value).astimezone(timezone(timedelta(hours=8))).replace(tzinfo=None)


class Reports:
    def __init__(self, store, settings):
        self.store, self.settings = store, settings
        self.lock = threading.Lock()
        self.path = settings.data_dir / "reports" / "orders.xlsx"

    def export(self):
        with self.lock:
            # 一次读事务保证明细、汇总与核对页取自同一数据库快照。
            with self.store.read() as db:
                db.execute("BEGIN")
                customers = {r["id"]: r["qq_number"] for r in db.execute("SELECT id,qq_number FROM customers")}
                orders = [dict(r) for r in db.execute("SELECT * FROM orders ORDER BY customer_id,id")]
                totals = [dict(r) for r in db.execute("SELECT c.id customer_id,COUNT(o.id) order_count,COALESCE(SUM(o.paid_cents),0) paid_cents,"
                           "COALESCE(SUM(o.refund_cents),0) refund_cents,COALESCE(SUM(o.net_cents),0) net_cents FROM customers c "
                           "LEFT JOIN orders o ON o.customer_id=c.id GROUP BY c.id")]
                reviews = [dict(r) for r in db.execute("SELECT * FROM reviews WHERE state='pending'")]
                checks = [dict(r) for r in db.execute("SELECT o.*,r.reason FROM review_checks r JOIN orders o ON o.id=r.order_id")]
                notices = [dict(r) for r in db.execute("SELECT * FROM notifications WHERE state NOT IN ('sent','simulated')")]
            wb = Workbook()
            wb.remove(wb.active)
            local_time = datetime.now(timezone(timedelta(hours=8))).replace(tzinfo=None)
            note = f"业务快照生成于 {local_time:%Y-%m-%d %H:%M}（北京时间）。累计有效消费按实付款减已到账退款；待付款、已取消、已退款计零。"
            def qq(cid):
                return customers.get(cid)

            self.sheet(wb, "客户汇总", ["QQ号", "订单数", "实付款合计（元）", "退款合计（元）", "累计有效消费（元）"],
                       [[qq(r["customer_id"]), r["order_count"], r["paid_cents"] / 100, r["refund_cents"] / 100, r["net_cents"] / 100] for r in totals], note + " QQ号由客户提供，空白表示尚未登记。",
                       {3, 4, 5})
            self.sheet(wb, "订单明细", ["QQ号", "淘宝账号名", "店铺", "订单号", "下单日期", "下单金额（元）", "实付款（元）", "累计退款（元）", "订单状态", "有效消费（元）", "更新来源", "最后更新", "权威核验时间"],
                       [[qq(r["customer_id"]), r["taobao_account"], r["shop"], r["order_no"], date.fromisoformat(r["ordered_on"]), r["order_cents"] / 100,
                         r["paid_cents"] / 100, r["refund_cents"] / 100, r["status"], r["net_cents"] / 100, {"customer_confirmation": "客户确认", "qq_text": "客户文字", "commerce_api": "店铺接口", "store_csv": "店铺导入"}.get(r["source"], "管理员/历史记录"), local_datetime(r["updated_at"]), local_datetime(r["verified_at"])] for r in orders],
                       "本页为数据库快照。请通过后台更正订单；直接编辑本文件不会回写数据库。订单号按文本保存。", {6, 7, 8, 10})
            pending_rows = []
            for r in reviews:
                draft = json.loads(r["draft"])
                effective = dict(draft)
                matched = [o for o in orders if o["customer_id"] == r["customer_id"] and o["order_no"] == draft.get("order_no")
                           and (not draft.get("shop") or o["shop"] == draft["shop"])]
                if len(matched) == 1:
                    old = matched[0]
                    effective = {**{k: old[k] for k in ("order_no", "shop", "taobao_account", "ordered_on", "status")},
                                 "order_amount": old["order_cents"] / 100, "paid_amount": old["paid_cents"] / 100,
                                 "refund_amount": old["refund_cents"] / 100, **{k: v for k, v in draft.items() if v is not None}}
                missing = [FIELD_NAMES[k] for k in REQUIRED if effective.get(k) is None or effective.get(k) == ""]
                if not qq(r["customer_id"]):
                    missing.insert(0, "QQ号")
                if effective.get("status") in ("部分退款", "退款中") and effective.get("refund_amount") is None:
                    missing.append(FIELD_NAMES["refund_amount"])
                def amount(name):
                    value = effective.get(name)
                    try:
                        if value is None:
                            return None
                        numeric = float(value)
                        return numeric if math.isfinite(numeric) else str(value)
                    except (ValueError, TypeError):
                        return str(value)
                day = effective.get("ordered_on")
                try:
                    day = date.fromisoformat(day) if day else None
                except (ValueError, TypeError):
                    pass
                pending_rows.append([r["id"], qq(r["customer_id"]), effective.get("taobao_account"), effective.get("shop"),
                                     effective.get("order_no"), day, amount("order_amount"), amount("paid_amount"),
                                     effective.get("status"), amount("refund_amount"), "、".join(missing), friendly_reason(r["reason"]), local_datetime(r["created_at"])])
            self.sheet(wb, "待确认识别", ["提交编号", "QQ号", "淘宝账号名", "店铺", "订单号", "下单日期", "下单金额（元）", "实付款（元）", "订单状态", "累计退款（元）", "待补充信息", "处理提示", "创建时间"],
                       pending_rows, "缺失信息保留空白，不计入消费。/补充 编号 补齐资料，/更正 编号 修改资料，核对后 /确认 编号。", {7, 8, 10}, wrap={11, 12})
            self.sheet(wb, "半月核对", ["QQ号", "淘宝账号名", "订单号", "店铺", "当前状态", "待核对原因", "权威核验时间"],
                       [[qq(r["customer_id"]), r["taobao_account"], r["order_no"], r["shop"], r["status"], r["reason"], local_datetime(r["verified_at"])] for r in checks],
                       f"每 {self.settings.reconcile_days} 天检查全部订单。没有权威数据时保持待核对，不会假称订单状态已刷新。", wrap={6})
            self.sheet(wb, "待发送通知", ["通知编号", "QQ号", "话术", "发送状态", "失败原因", "尝试次数", "创建时间"],
                       [[r["id"], qq(r["customer_id"]), r["content"], r["state"], r["error"], r["attempts"], local_datetime(r["created_at"])] for r in notices],
                       "waiting 等待可用消息窗口；failed 需处理权限或网络问题；QQ_DRY_RUN 模式仅模拟发送。", wrap={3, 5})
            self.path.parent.mkdir(parents=True, exist_ok=True)
            temp = self.path.with_name(f".{uuid.uuid4().hex}.xlsx")
            try:
                wb.save(temp)
                os.replace(temp, self.path)
            finally:
                if temp.exists():
                    temp.unlink()
            return self.path

    @staticmethod
    def sheet(wb, name, headers, rows, note, money=frozenset(), wrap=frozenset()):
        ws = wb.create_sheet(name)
        width = len(headers)
        ws.append([name])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        ws["A1"].font = Font(name="Arial", size=16, bold=True, color="17324D")
        ws.row_dimensions[1].height = 32
        ws.append([note])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
        ws["A2"].font = Font(name="Arial", size=10, color="50657A")
        ws.row_dimensions[2].height = 40
        ws.append(headers)
        for cell in ws[3]:
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.font = Font(name="Arial", color="FFFFFF", bold=True, size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[3].height = 30
        for index, row in enumerate(rows, 4):
            for col, value in enumerate(row, 1):
                cell = ws.cell(index, col, value)
                # 用户输入不得作为公式执行，包括 =、+、-、@ 开头的店铺等文本。
                if isinstance(value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                cell.font = Font(name="Arial", size=11, color="17324D")
                cell.alignment = Alignment(horizontal="left" if isinstance(value, str) else "right", vertical="top" if col in wrap else "center", wrap_text=col in wrap)
                if col in money:
                    cell.number_format = '#,##0.00'
                elif isinstance(value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                elif isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
                if index % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EFF4F8")
            lines = max((sum(max(1, math.ceil(sum(2 if ord(ch) > 127 else 1 for ch in line) / 46)) for line in str(row[col - 1] or "").splitlines()) for col in wrap), default=1)
            ws.row_dimensions[index].height = min(409, max(24, 16 * lines + 8))
        for col, header in enumerate(headers, 1):
            column_width = 48 if col in wrap else 38 if "时间" in header or "更新" in header else 36 if "ID" in header or "事件" in header else 24
            ws.column_dimensions[get_column_letter(col)].width = column_width
        ws.freeze_panes = "C4" if width > 3 else "A4"
        ws.auto_filter.ref = f"A3:{get_column_letter(width)}{max(3, ws.max_row)}"
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:3"
