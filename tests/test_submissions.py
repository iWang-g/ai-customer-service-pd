import asyncio
import json
import sqlite3

import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from conftest import event
from qqbot.app import create_app
from qqbot.extraction import Draft, parse_text
from qqbot.storage import Store


def send(app, text, user="u1", message=None):
    message = message or str(app.state.store.rows("SELECT COUNT(*) n FROM notifications")[0]["n"])
    asyncio.run(app.state.service.process("event-" + message, event(message, user=user, content=text)))


def template(number="000123456789", account="买家甲", qq="123456789"):
    return f"QQ号：{qq}\n淘宝账号名：{account}\n店铺：示例店\n订单号：{number}\n下单日期：2026-10-13\n下单金额：120.00\n实付款：99.00\n订单状态：待收货"


def latest_reply(app):
    return app.state.store.rows("SELECT content FROM notifications ORDER BY id DESC LIMIT 1")[0]["content"]


def test_template_order_and_partial_labels():
    data = parse_text(template())[0].data
    assert data["shop"] == "示例店"
    assert data["qq_number"] == "123456789"
    assert data["taobao_account"] == "买家甲"
    assert data["order_no"] == "000123456789"
    assert parse_text("QQ号：\n淘宝账号名：买家甲\n实付款：0")[0].data == {"taobao_account": "买家甲", "paid_amount": "0"}
    records = parse_text(template() + "\n\n" + template("000123456790", "买家乙"))
    assert len(records) == 2
    assert records[0].data["taobao_account"] == "买家甲"
    assert records[1].data["taobao_account"] == "买家乙"


def test_complete_text_and_qq_reuse_per_order_account(settings):
    app = create_app(settings)
    with TestClient(app):
        send(app, template())
        send(app, template("000123456790", "买家乙", ""))
        orders = app.state.store.rows("SELECT * FROM orders ORDER BY id")
        assert [o["taobao_account"] for o in orders] == ["买家甲", "买家乙"]
        assert app.state.store.totals()[0]["net_cents"] == 19800
        assert app.state.store.totals()[0]["qq_number"] == "123456789"
        send(app, template("000123456790", "买家乙", ""))
        assert len(app.state.store.rows("SELECT * FROM orders")) == 2
        assert "没有重复" in latest_reply(app)
        send(app, "订单号：000123456790\n店铺：示例店\n订单状态：部分退款\n退款金额：20.00")
        assert app.state.store.totals()[0]["net_cents"] == 17800
        assert app.state.store.rows("SELECT taobao_account FROM orders WHERE order_no='000123456790'")[0]["taobao_account"] == "买家乙"


def test_missing_order_number_supplement_restart_confirm_once(settings):
    app = create_app(settings)
    with TestClient(app):
        send(app, "店铺：示例店\n实付款：99.00\n订单状态：待收货", message="initial")
        assert not app.state.store.rows("SELECT * FROM orders")
        row = app.state.store.rows("SELECT * FROM reviews")[0]
        number = row["id"]
        assert "order_no" not in json.loads(row["draft"])
        assert "还差" in latest_reply(app) and "order_no" not in latest_reply(app)
        send(app, f"/确认 {number}")
        assert not app.state.store.rows("SELECT * FROM orders")
    app = create_app(settings)
    with TestClient(app):
        send(app, "QQ号：123456789\n淘宝账号名：买家甲\n订单号：000123456789\n下单日期：2026-10-13\n下单金额：120.00", message="supplement")
        rows = app.state.store.rows("SELECT * FROM reviews")
        assert len(rows) == 1 and rows[0]["id"] == number
        assert not app.state.store.rows("SELECT * FROM orders")
        send(app, f"/确认 {number}", message="confirm")
        send(app, f"/确认 {number}", message="confirm-twice")
        assert len(app.state.store.rows("SELECT * FROM orders")) == 1
        assert app.state.store.totals()[0]["net_cents"] == 9900
        assert app.state.store.rows("SELECT state FROM reviews")[0]["state"] == "accepted"
        # Crash/retry after a committed supplement must not recreate a pending record.
        send(app, "QQ号：123456789\n淘宝账号名：买家甲\n订单号：000123456789\n下单日期：2026-10-13\n下单金额：120.00", message="supplement")
        assert not app.state.store.rows("SELECT * FROM reviews WHERE state='pending'")


def test_multiple_pending_conflicts_and_owner(settings):
    app = create_app(settings)
    with TestClient(app):
        send(app, "订单号：aaa\n店铺：示例店\n实付款：99.00")
        send(app, "订单号：bbb\n店铺：示例店\n实付款：88.00")
        send(app, "淘宝账号名：买家甲")
        assert "/补充 编号" in latest_reply(app)
        rows = app.state.store.rows("SELECT * FROM reviews ORDER BY id")
        assert len(rows) == 2
        number = rows[0]["id"]
        send(app, f"/补充 {number}\n实付款：90.00")
        assert "不一致" in latest_reply(app)
        assert json.loads(app.state.store.rows("SELECT draft FROM reviews WHERE id=?", (number,))[0]["draft"])["paid_amount"] == "99.00"
        send(app, f"/更正 {number}\n实付款：90.00")
        assert json.loads(app.state.store.rows("SELECT draft FROM reviews WHERE id=?", (number,))[0]["draft"])["paid_amount"] == "90.00"
        send(app, f"/补充 {number}\n淘宝账号名：陌生人", user="another")
        assert "属于您的" in latest_reply(app)
        assert "taobao_account" not in json.loads(app.state.store.rows("SELECT draft FROM reviews WHERE id=?", (number,))[0]["draft"])


def test_qq_does_not_merge_or_silently_change(settings):
    app = create_app(settings)
    with TestClient(app):
        send(app, "QQ号：123456789")
        send(app, "QQ号：123456789", user="another")
        assert len(app.state.store.rows("SELECT * FROM customers")) == 2
        send(app, "QQ号：987654321")
        assert "不一致" in latest_reply(app)
        assert {r["qq_number"] for r in app.state.store.rows("SELECT qq_number FROM customers")} == {"123456789"}


def test_invalid_fields_saved_and_explicitly_corrected(settings):
    app = create_app(settings)
    with TestClient(app):
        send(app, "QQ号：xxx\n订单号：abc\n下单日期：不是日期\n实付款：NaN")
        row = app.state.store.rows("SELECT * FROM reviews")[0]
        assert "QQ号" in latest_reply(app) and "格式" in latest_reply(app)
        assert not app.state.store.rows("SELECT * FROM orders")
        app.state.service.reports.export()  # Malformed draft values must still be exportable.
        send(app, f"/更正 {row['id']}\n" + template("abc"))
        send(app, f"/确认 {row['id']}")
        assert app.state.store.totals()[0]["net_cents"] == 9900


def test_unknown_refund_is_not_zero_until_supplied(settings):
    app = create_app(settings)
    with TestClient(app):
        send(app, template().replace("待收货", "退款中"))
        row = app.state.store.rows("SELECT * FROM reviews")[0]
        assert "refund_amount" not in json.loads(row["draft"])
        assert "累计已到账退款：" in latest_reply(app)
        send(app, f"/补充 {row['id']}\n退款金额：0")
        send(app, f"/确认 {row['id']}")
        assert app.state.store.totals()[0]["net_cents"] == 9900


def test_image_partial_preserves_evidence_through_supplement(settings, monkeypatch):
    app = create_app(settings)
    with TestClient(app):
        async def fake_extract(text, attachments):
            app.state.service.extractor.evidence = ["image-digest"]
            return [Draft({"shop": "示例店", "paid_amount": "99.00", "status": "待收货"}, .99, "ai")]
        original = app.state.service.extractor.extract
        monkeypatch.setattr(app.state.service.extractor, "extract", fake_extract)
        send(app, "", message="image")
        number = app.state.store.rows("SELECT id FROM reviews")[0]["id"]
        monkeypatch.setattr(app.state.service.extractor, "extract", original)
        send(app, f"/补充 {number}\nQQ号：123456789\n淘宝账号名：买家甲\n订单号：abc\n下单日期：2026-10-13\n下单金额：120.00")
        assert app.state.store.rows("SELECT * FROM evidence")[0]["event_id"] == "event-image"
        assert not app.state.store.rows("SELECT * FROM orders")
        send(app, f"/确认 {number}")
        assert app.state.store.totals()[0]["net_cents"] == 9900


def test_missing_images_in_batch_are_not_deduplicated(settings):
    from qqbot.extraction import Extractor
    settings.model_base_url = "https://model.test"
    settings.model_api_key = "test"
    settings.vision_model = "vision"
    async def run():
        response = {"orders": [{"shop": "店甲", "confidence": .5}, {"shop": "店甲", "confidence": .5}]}
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(200, json={"choices": [{"message": {"content": json.dumps(response)}}]}))) as http:
            drafts = await Extractor(settings, http).extract("订单信息不完整", [])
            assert len(drafts) == 2
    asyncio.run(run())


def test_report_columns_missing_values_and_text_safety(settings):
    app = create_app(settings)
    with TestClient(app):
        send(app, template(account="=1+1"))
        send(app, "淘宝账号名：买家乙\n订单号：other\n店铺：另一店")
        path = app.state.service.reports.export()
        wb = load_workbook(path)
        ws = wb["订单明细"]
        assert [c.value for c in ws[3]][:4] == ["QQ号", "淘宝账号名", "店铺", "订单号"]
        assert ws["A4"].value == "123456789" and ws["A4"].data_type == "s"
        assert ws["B4"].value == "=1+1" and ws["B4"].data_type == "s"
        assert ws["D4"].value == "000123456789"
        pending = wb["待确认识别"]
        assert pending["G4"].value is None and pending["J4"].value is None
        assert "下单日期" in pending["K4"].value
        assert not any("昵称" in str(c.value) or "客户ID" == c.value for sheet in wb for c in sheet[3])
        wb.close()


def test_old_database_migration_is_repeatable(tmp_path):
    path = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(path) as db:
        db.execute("CREATE TABLE customers(id TEXT PRIMARY KEY,created_at TEXT NOT NULL)")
        db.execute("INSERT INTO customers VALUES ('old','2026-10-13')")
    Store(path)
    store = Store(path)
    assert store.rows("SELECT * FROM customers")[0]["id"] == "old"
    assert store.rows("SELECT * FROM customers")[0]["qq_number"] is None


def test_admin_customer_edit_is_authenticated_and_audited(settings):
    app = create_app(settings)
    with TestClient(app) as client:
        send(app, "QQ号：123456789")
        cid = app.state.store.totals()[0]["customer_id"]
        body = {"qq_number": "987654321", "verification_note": "已核实修改"}
        assert client.patch(f"/admin/customers/{cid}", json=body).status_code == 401
        response = client.patch(f"/admin/customers/{cid}", json=body, headers={"Authorization": "Bearer " + settings.admin_token})
        assert response.status_code == 200
        assert app.state.store.totals()[0]["qq_number"] == "987654321"
        assert len(app.state.store.rows("SELECT * FROM audit WHERE action='customer_qq'")) == 2
