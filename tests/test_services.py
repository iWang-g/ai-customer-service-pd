import asyncio
import io
import json
import time

import httpx
import pytest
from openpyxl import load_workbook
from PIL import Image

from conftest import event
from qqbot.app import create_app
from qqbot.domain import OrderInput
from qqbot.extraction import ExtractionError, Extractor, parse_text
from qqbot.notifications import Notifier
from qqbot.qq import QQClient, SendError
from qqbot.reconcile import Reconciler
from qqbot.reports import Reports


def test_text_template_keeps_long_identifier_and_exact_amount():
    text = "订单号：00012345678901234567\n店铺：测试旗舰店\n下单日期：2026年10月1日\n下单金额：￥120.00\n实付款：99.90元\n订单状态：待收货"
    order = OrderInput(**parse_text(text)[0].data)
    assert order.order_no == "00012345678901234567"
    assert str(order.paid_amount) == "99.9"
    assert order.ordered_on.isoformat() == "2026-10-01"


def test_duplicate_order_in_same_text_requires_confirmation():
    drafts = parse_text("订单号：123 店铺：商店 实付款：10\n订单号：123 店铺：商店 实付款：20")
    assert len(drafts) == 2 and all(d.reason for d in drafts)


def test_no_interface_means_unverified_not_fake_refresh(ledger, cid, full_order, settings):
    ledger.apply(cid, OrderInput(**full_order), "test")
    async def run():
        async with httpx.AsyncClient() as http:
            return await Reconciler(ledger.store, ledger, settings, http).run()
    result = asyncio.run(run())
    assert result["needs_review"] == 1
    assert result["verified"] == 0
    assert ledger.store.rows("SELECT verified_at FROM orders")[0]["verified_at"] is None
    assert ledger.store.get_meta("last_full_reconcile_completed") is None


def test_commerce_adapter_checks_identity_and_refunds(ledger, cid, full_order, settings):
    ledger.apply(cid, OrderInput(**full_order), "test")
    settings.commerce_base_url = "https://commerce.example"
    async def run(data):
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(200, json=data))) as http:
            return await Reconciler(ledger.store, ledger, settings, http).run()
    assert asyncio.run(run({**full_order, "order_no": "wrong", "refund_amount": "99.90", "status": "已退款"}))["needs_review"] == 1
    assert ledger.store.totals()[0]["net_cents"] == 9990
    assert asyncio.run(run({**full_order, "refund_amount": "99.90", "status": "已退款"}))["verified"] == 1
    assert ledger.store.totals()[0]["net_cents"] == 0
    assert not ledger.store.rows("SELECT * FROM review_checks")


def test_waiting_notification_wakes_only_with_new_message(ledger, cid, settings):
    with ledger.store.transaction() as db:
        db.execute("UPDATE identities SET last_event_at=?", (time.time() - 4000,))
    ledger.store.notify("test", cid, "消费有变更")
    async def run():
        async with httpx.AsyncClient() as http:
            notifier = Notifier(ledger.store, QQClient(settings, http), settings)
            await notifier.flush_one()
            assert ledger.store.rows("SELECT state FROM notifications")[0]["state"] == "waiting"
            ledger.store.identity(event("m2"))
            notifier.wake_customer(cid)
            await notifier.flush_one()
    asyncio.run(run())
    assert ledger.store.rows("SELECT state FROM notifications")[0]["state"] == "simulated"


def test_five_reply_limit_keeps_extra_notifications_waiting(ledger, cid, settings):
    for number in range(6):
        ledger.store.notify(str(number), cid, "notice")
    async def run():
        async with httpx.AsyncClient() as http:
            notifier = Notifier(ledger.store, QQClient(settings, http), settings)
            for _ in range(6):
                await notifier.flush_one()
    asyncio.run(run())
    states = [r["state"] for r in ledger.store.rows("SELECT state FROM notifications ORDER BY id")]
    assert states == ["simulated"] * 5 + ["waiting"]


def test_send_timeout_reuses_message_sequence(ledger, cid, settings):
    attempts = []
    class FlakyQQ:
        async def send(self, identity, content, msg_id, seq):
            attempts.append((msg_id, seq))
            if len(attempts) == 1:
                raise SendError("rate limit", retryable=True)
            return "sent"
    ledger.store.notify("test", cid, "notice")
    async def run():
        notifier = Notifier(ledger.store, FlakyQQ(), settings)
        await notifier.flush_one()
        with ledger.store.transaction() as db:
            db.execute("UPDATE notifications SET available_at=0")
        await notifier.flush_one()
    asyncio.run(run())
    assert attempts == [("m1", 1), ("m1", 1)]


def test_qq_token_refresh_and_request_contract(settings):
    settings.qq_dry_run = False
    calls = []
    def handler(request):
        calls.append(request)
        if request.url.path.endswith("getAppAccessToken"):
            return httpx.Response(200, json={"access_token": "token", "expires_in": "7200"})
        assert request.headers["Authorization"] == "QQBot token"
        assert request.headers["X-Union-Appid"] == "test-app"
        assert json.loads(request.content)["msg_id"] == "m1"
        return httpx.Response(401 if len(calls) == 2 else 200, json={"id": "sent"})
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            qq = QQClient(settings, http)
            assert await qq.send({"scene": "c2c", "target_id": "u1"}, "hello", "m1", 1) == "sent"
            with pytest.raises(SendError, match="主动"):
                await qq.send({"scene": "c2c", "target_id": "u1"}, "hello")
    asyncio.run(run())
    assert len(calls) == 4


def test_image_whitelist_prevents_internal_download(settings):
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r: pytest.fail("不应发出网络请求"))) as http:
            with pytest.raises(ExtractionError, match="白名单"):
                await Extractor(settings, http).image_data("http://127.0.0.1/secret")
    asyncio.run(run())


def test_current_qq_image_cdn_is_allowed(settings):
    image = io.BytesIO()
    Image.new("RGB", (8, 8), "white").save(image, "PNG")

    async def run():
        async with httpx.AsyncClient(
            transport=httpx.MockTransport(lambda request: httpx.Response(200, content=image.getvalue()))
        ) as http:
            result = await Extractor(settings, http).image_data(
                "https://multimedia.nt.qq.com.cn/example.jpg"
            )
            assert result["image_url"]["url"].startswith("data:image/png;base64,")

    asyncio.run(run())


def test_image_extraction_pending_confirmation_with_evidence(settings, full_order):
    from fastapi.testclient import TestClient
    settings.model_base_url = "https://model.example/v1"
    settings.model_api_key = "test-model-key"
    settings.vision_model = "vision"
    image = io.BytesIO()
    Image.new("RGB", (32, 32), "white").save(image, "PNG")
    def handler(request):
        if request.url.host == "gchat.qpic.cn":
            return httpx.Response(200, content=image.getvalue())
        body = json.loads(request.content)
        assert body["messages"][1]["content"][1]["image_url"]["url"].startswith("data:image/png;base64,")
        return httpx.Response(200, json={"choices": [{"message": {"content": json.dumps({"orders": [{**full_order, "confidence": .99, "uncertainties": []}]})}}]})
    app = create_app(settings, httpx.MockTransport(handler))
    with TestClient(app):
        async def run():
            await app.state.service.process("image-event", event(attachments=[{"url": "https://gchat.qpic.cn/img.png", "content_type": "image/png"}]))
            assert not app.state.store.rows("SELECT * FROM orders")
            review = app.state.store.rows("SELECT * FROM reviews")[0]
            assert review["state"] == "pending"
            assert len(app.state.store.rows("SELECT * FROM evidence")) == 1
            await app.state.service.process("supplement-event", event("m-supplement", content=f"/补充 {review['id']}\nQQ号：123456789\n淘宝账号名：测试买家"))
            await app.state.service.process("confirm-event", event("m2", content=f"/确认 {review['id']}"))
        asyncio.run(run())
        assert app.state.store.totals()[0]["net_cents"] == 9990


def test_export_keeps_ids_numeric_money_and_text_safe(ledger, cid, full_order, settings):
    ledger.apply(cid, OrderInput(**{**full_order, "shop": "=1+1"}), "test")
    path = Reports(ledger.store, settings).export()
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ["客户汇总", "订单明细", "待确认识别", "半月核对", "待发送通知"]
    ws = wb["订单明细"]
    assert ws["D4"].value == full_order["order_no"]
    assert ws["D4"].data_type == "s"
    assert ws["D4"].number_format == "@"
    assert ws["C4"].value == "=1+1" and ws["C4"].data_type == "s"
    assert ws["G4"].value == 99.9
    assert wb["客户汇总"]["E4"].value == 99.9
    assert ws.freeze_panes == "C4"
    wb.close()


def test_malformed_token_becomes_recorded_notification_failure(ledger, cid, settings):
    settings.qq_dry_run = False
    ledger.store.notify("test", cid, "notice")
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(200, text="invalid"))) as http:
            notifier = Notifier(ledger.store, QQClient(settings, http), settings)
            await notifier.flush_one()
    asyncio.run(run())
    assert ledger.store.rows("SELECT state FROM notifications")[0]["state"] == "failed"


def test_restart_recovers_queued_work(ledger):
    ledger.store.enqueue_event("e", event())
    assert ledger.store.claim_event()["id"] == "e"
    ledger.store.recover()
    assert ledger.store.claim_event()["id"] == "e"
